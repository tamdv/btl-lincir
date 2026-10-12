"""[NHOM8] Put the experiment results into the hand-edited deck, in place (Morph names and manual edits kept).

Inputs: results/assets.json, results/figures/*.png and results/analysis_vi.json (optional keys: title_14, title_15,
title_16, slide_sig [list], slide_conclusion [list]). A timestamped backup is written first.

  python slides/fill_results.py [--deck slides/Slides_Nhom8_LinCIR.pptx] [--dry_run]
"""
import argparse
import copy
import datetime as dt
import json
import re
import shutil
import statistics
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "results"
NAVY = RGBColor(0x0F, 0x1C, 0x3F)
VARIANTS = {"Baseline": "EXP-001-baseline", "A:": "EXP-002-infonce", "B:": "EXP-003-gapnoise", "A + B": "EXP-004-infonce-gapnoise"}
COLS = {"FIQ TB R@10": "average_fiq_recall_at10", "FIQ TB R@50": "average_fiq_recall_at50", "CIRCO mAP@5": "circo_map_at5"}


def vn(x):
    return re.sub(r"(\d)\.(\d)", r"\1,\2", x)


def ms(v):
    v = [x for x in (v or []) if x is not None]
    if not v:
        return None
    return vn(f"{v[0]:.2f}" if len(v) == 1 else f"{statistics.mean(v):.2f} ± {statistics.stdev(v):.2f}")


def set_text(tf, lines, color=NAVY):
    """replace a text frame's text, keeping the first run's formatting for every line"""
    lines = lines if isinstance(lines, list) else [lines]
    p0 = tf.paragraphs[0]
    r0 = p0.runs[0] if p0.runs else p0.add_run()
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    r0.text = lines[0]
    r0.font.color.rgb = color
    for line in lines[1:]:
        new_p = copy.deepcopy(p0._p)
        tf._txBody.append(new_p)
        tf.paragraphs[-1].runs[0].text = line
    return tf


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(ROOT / "slides" / "Slides_Nhom8_LinCIR.pptx"))
    ap.add_argument("--assets", default=str(RES / "assets.json"))
    ap.add_argument("--analysis", default=str(RES / "analysis_vi.json"))
    ap.add_argument("--figures", default=str(RES / "figures"))
    ap.add_argument("--dry_run", action="store_true")
    a = ap.parse_args()
    A = json.loads(Path(a.assets).read_text())
    T = json.loads(Path(a.analysis).read_text()) if Path(a.analysis).exists() else {}
    main_r = A.get("main", {})
    deck = Path(a.deck)
    prs = Presentation(deck)
    done, missing = [], []

    for idx, s in enumerate(prs.slides, start=1):
        title = next((sh for sh in s.shapes if sh.name == "!!title"), None)
        for sh in list(s.shapes):
            # results table
            if sh.has_table and "FIQ TB R@10" in [c.text for c in sh.table.rows[0].cells]:
                hdr = [c.text.strip() for c in sh.table.rows[0].cells]
                for row in list(sh.table.rows)[1:]:
                    label = row.cells[0].text.strip()
                    m = main_r.get(next((v for k, v in VARIANTS.items() if label.startswith(k)), ""), {})
                    for j, h in enumerate(hdr[1:], start=1):
                        if h in COLS and ms(m.get(COLS[h])):
                            set_text(row.cells[j].text_frame, ms(m[COLS[h]]))
                done.append(f"slide {idx}: bảng kết quả")
                if T.get("title_14") and title is not None:
                    set_text(title.text_frame, T["title_14"])
                    done.append(f"slide {idx}: tiêu đề")
                continue
            if not sh.has_text_frame:
                continue
            txt = sh.text_frame.text.strip()
            # improvement A stat (slide with the InfoNCE title)
            if txt == "[chờ]" and main_r.get("EXP-002-infonce"):
                set_text(sh.text_frame, ms(main_r["EXP-002-infonce"]["average_fiq_recall_at10"]).split(" ±")[0], color=RGBColor(0x1E, 0x6F, 0xE8))
                done.append(f"slide {idx}: số của A")
            elif txt.startswith("FashionIQ TB R@10 của A (baseline [chờ])") and main_r.get("EXP-001-baseline"):
                base = ms(main_r["EXP-001-baseline"]["average_fiq_recall_at10"]).split(" ±")[0]
                set_text(sh.text_frame, f"FashionIQ TB R@10 của A (baseline {base})", color=RGBColor(0x5B, 0x6B, 0x82))
                done.append(f"slide {idx}: nhãn baseline")
            elif txt.startswith("[Chờ kết quả phiên eval"):
                lines = T.get("slide_sig")
                if lines:
                    set_text(sh.text_frame, lines)
                    done.append(f"slide {idx}: dòng kiểm định")
                else:
                    missing.append(f"slide {idx}: slide_sig")
            elif txt.startswith("[Chờ kết quả:") and txt.endswith(".png]"):
                fig = Path(a.figures) / txt[len("[Chờ kết quả:"):-1].strip()
                if fig.with_name(fig.stem + "_slide.png").exists():  # larger-text version for slides
                    fig = fig.with_name(fig.stem + "_slide.png")
                frame = next((o for o in s.shapes if o.shape_type == 1 and o.left <= sh.left and o.top <= sh.top
                              and o.left + o.width >= sh.left + sh.width and o.name.startswith("Rounded")), None)
                if not fig.exists():
                    missing.append(f"slide {idx}: {fig.name}")
                    continue
                box = frame or sh
                from PIL import Image
                iw, ih = Image.open(fig).size
                bw, bh = box.width, max(box.height, Emu(int(4.6 * 914400)))  # use the free space down to the footer
                scale = min(bw / iw, bh / ih)
                w, h = int(iw * scale), int(ih * scale)
                s.shapes.add_picture(str(fig), box.left + (bw - w) // 2, box.top, w, h)
                for o in (sh, frame):
                    if o is not None:
                        o._element.getparent().remove(o._element)
                done.append(f"slide {idx}: hình {fig.name}")
                if T.get("title_15") and title is not None and fig.name.startswith("modality_gap"):
                    set_text(title.text_frame, T["title_15"])
            elif txt.startswith("[Kết luận H1–H3"):
                if T.get("slide_conclusion"):
                    set_text(sh.text_frame, T["slide_conclusion"])
                    done.append(f"slide {idx}: kết luận")
                    if T.get("title_16") and title is not None:
                        set_text(title.text_frame, T["title_16"])
                else:
                    missing.append(f"slide {idx}: slide_conclusion")

    print("Đã điền:", *done, sep="\n  - ")
    if missing:
        print("Chưa điền được:", *missing, sep="\n  - ")
    if a.dry_run:
        return
    backup = deck.with_name(f"{deck.stem}_truocKhiDienKQ_{dt.datetime.now():%Y%m%d-%H%M}.pptx")
    shutil.copy(deck, backup)
    prs.save(deck)
    print("Sao lưu:", backup.name, "| đã ghi:", deck.name)


if __name__ == "__main__":
    main()
