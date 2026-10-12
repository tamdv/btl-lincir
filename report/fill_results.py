"""[NHOM8] Put the experiment results into the hand-edited report, in place, without touching the rest.

Inputs: results/assets.json + results/figures/*.png (scripts/make_report_assets.py) and results/analysis_vi.json
(the written analysis: keys main, sig, gap, curve, qual, conclusion, speed). A timestamped backup is written first.

  python report/fill_results.py [--doc report/BaoCao_Nhom8_LinCIR.docx] [--dry_run]

Finds its targets by content (table headers, "[CHỜ KẾT QUẢ] …" paragraphs), so it keeps working after manual edits;
prints every target it could not find instead of guessing.
"""
import argparse
import copy
import datetime as dt
import json
import re
import shutil
import statistics
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "results"
VARIANTS = {"Baseline": "EXP-001-baseline", "A:": "EXP-002-infonce", "B:": "EXP-003-gapnoise", "A + B": "EXP-004-infonce-gapnoise"}
HEADER_KEYS = {"Dress R@10": "dress_fiq_recall_at10", "Shirt R@10": "shirt_fiq_recall_at10", "Toptee R@10": "toptee_fiq_recall_at10",
               "FIQ TB R@10": "average_fiq_recall_at10", "FIQ TB R@50": "average_fiq_recall_at50",
               "CIRCO mAP@5": "circo_map_at5", "CIRCO mAP@10": "circo_map_at10"}
# placeholder paragraph (start of its text after "[CHỜ KẾT QUẢ] ") -> key in analysis_vi.json, or a figure file
TEXT_SLOTS = {"Nhận xét bảng kết quả chính": "main", "Kết luận H1–H3 dựa trên": "sig", "Phân tích: biến thể nào làm": "gap",
              "Nhận xét đường học": "curve", "Bình luận ví dụ": "qual", "Tóm tắt kết quả chính": "conclusion"}
FIGURE_SLOTS = {"hình modality_gap.png": "modality_gap.png", "hình learning_curve.png": "learning_curve.png",
                "hình qualitative_fiq_1.png": "qualitative_fiq_1.png"}


def vn(x):
    return re.sub(r"(\d)\.(\d)", r"\1,\2", x)


def ms(v):
    v = [x for x in (v or []) if x is not None]
    if not v:
        return None
    return vn(f"{v[0]:.2f}" if len(v) == 1 else f"{statistics.mean(v):.2f} ± {statistics.stdev(v):.2f}")


def set_cell(cell, value):
    p = cell.paragraphs[0]
    runs = p.runs
    if runs:
        runs[0].text = value
        runs[0].font.highlight_color = None
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(value)


def set_paragraph(p, text, justify=True):
    """replace the paragraph's text with plain runs (keeps its style); **bold** supported"""
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part:
            r = p.add_run(part[2:-2] if part.startswith("**") else part)
            r.bold = part.startswith("**") or None
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--doc", default=str(ROOT / "report" / "BaoCao_Nhom8_LinCIR.docx"))
    ap.add_argument("--assets", default=str(RES / "assets.json"))
    ap.add_argument("--analysis", default=str(RES / "analysis_vi.json"))
    ap.add_argument("--figures", default=str(RES / "figures"))
    ap.add_argument("--dry_run", action="store_true")
    a = ap.parse_args()
    A = json.loads(Path(a.assets).read_text())
    T = json.loads(Path(a.analysis).read_text()) if Path(a.analysis).exists() else {}
    doc_path = Path(a.doc)
    d = docx.Document(doc_path)
    done, missing = [], []

    # main results table
    t_main = next((t for t in d.tables if t.rows[0].cells[0].text.strip() == "Biến thể" and "FIQ TB R@10" in [c.text.strip() for c in t.rows[0].cells]), None)
    if t_main is None:
        missing.append("bảng kết quả chính (Bảng 5.3)")
    else:
        hdr = [c.text.strip() for c in t_main.rows[0].cells]
        for row in t_main.rows[1:]:
            label = row.cells[0].text.strip()
            prefix = next((v for k, v in VARIANTS.items() if label.startswith(k)), None)
            m = A["main"].get(prefix) if prefix else None
            if not m:
                missing.append(f"số liệu cho dòng '{label}'")
                continue
            for j, h in enumerate(hdr[1:], start=1):
                if h in HEADER_KEYS and ms(m.get(HEADER_KEYS[h])):
                    set_cell(row.cells[j], ms(m[HEADER_KEYS[h]]))
        done.append("Bảng 5.3 (kết quả chính)")

    # significance table: one row per (variant, dataset, metric)
    t_sig = next((t for t in d.tables if t.rows[0].cells[0].text.strip() == "So với baseline"), None)
    if t_sig is None:
        missing.append("bảng kiểm định (Bảng 5.4)")
    else:
        rows = []
        for prefix, label in (("EXP-002-infonce", "A"), ("EXP-003-gapnoise", "B"), ("EXP-004-infonce-gapnoise", "A+B")):
            for r in A["significance"].get(prefix, []):
                if r["metric"] in ("R@10", "R@50", "ap5", "ap10"):
                    pw = "—" if r["p_wilcoxon"] is None else f"{r['p_wilcoxon']:.4f}"
                    rows.append([label, r["dataset"], r["metric"].replace("ap", "mAP@"), f"{r['delta']:+.2f}",
                                 f"[{r['ci95'][0]:+.2f}; {r['ci95'][1]:+.2f}]", f"{r['p_boot']:.4f}", pw])
        template = copy.deepcopy(t_sig.rows[1]._tr)
        for tr in [r._tr for r in t_sig.rows[1:]]:
            tr.getparent().remove(tr)
        for vals in rows:
            tr = copy.deepcopy(template)
            t_sig._tbl.append(tr)
            new_row = t_sig.rows[-1]
            for j, v in enumerate(vals):
                set_cell(new_row.cells[j], vn(v))
        done.append(f"Bảng 5.4 (kiểm định, {len(rows)} dòng)")

    # placeholder paragraphs: analysis text or figures
    for p in d.paragraphs:
        if not p.text.startswith("[CHỜ KẾT QUẢ]"):
            continue
        body = p.text[len("[CHỜ KẾT QUẢ]"):].strip()
        fig = next((f for k, f in FIGURE_SLOTS.items() if body.startswith(k)), None)
        key = next((v for k, v in TEXT_SLOTS.items() if body.startswith(k)), None)
        if fig:
            path = Path(a.figures) / fig
            if path.exists():
                for r in list(p.runs):
                    r._r.getparent().remove(r._r)
                p.add_run().add_picture(str(path), width=Cm(15))
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                done.append(f"hình {fig}")
            else:
                missing.append(f"hình {fig}")
        elif key:
            if T.get(key):
                set_paragraph(p, T[key])
                done.append(f"đoạn '{key}'")
            else:
                missing.append(f"chữ phân tích '{key}' trong {Path(a.analysis).name}")
        else:
            missing.append(f"chưa biết điền gì cho: {p.text[:60]}")

    # measured speed (fact known before the evaluation)
    for p in d.paragraphs:
        if "tốc độ khoảng 2,4 giây/bước" in p.text and T.get("speed"):
            set_paragraph(p, p.text.replace("tốc độ khoảng 2,4 giây/bước.", T["speed"]))
            done.append("câu tốc độ huấn luyện (5.1)")

    print("Đã điền:", *done, sep="\n  - ")
    if missing:
        print("Chưa điền được:", *missing, sep="\n  - ")
    if a.dry_run:
        return
    backup = doc_path.with_name(f"{doc_path.stem}_truocKhiDienKQ_{dt.datetime.now():%Y%m%d-%H%M}.docx")
    shutil.copy(doc_path, backup)
    d.save(doc_path)
    print("Sao lưu:", backup.name, "| đã ghi:", doc_path.name)


if __name__ == "__main__":
    main()
