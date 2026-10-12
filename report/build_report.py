"""[NHOM8] Build the report .docx from the group's template (cover, styles) + the content below.

Result numbers and figures come from results/assets.json and results/figures/ (scripts/make_report_assets.py).
Without them, the results sections show yellow placeholders. Rerun after every new result:

  python report/make_figures.py && python scripts/make_report_assets.py && python report/build_report.py

Then open the .docx in Word and update fields (Cmd+A, F9) to refresh the table of contents and lists of figures/tables.
"""
import copy
import json
import re
import shutil
import statistics
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path("/Users/tamdv/Documents/Workspace/Master/3.HĐHM/NOS – 2025-2 - Nhóm 10.docx")
OUT = ROOT / "report" / "BaoCao_Nhom8_LinCIR.docx"
FIG = ROOT / "report" / "figures"
RES_FIG = ROOT / "results" / "figures"
ASSETS = ROOT / "results" / "assets.json"

COVER = {
    "BÁO CÁO TIỂU LUẬN": "BÁO CÁO BÀI TẬP LỚN",
    "MÔN HỆ ĐIỀU HÀNH MẠNG": "MÔN TÌM KIẾM VÀ TRUY XUẤT THÔNG TIN",
    "TÌM HIỂU MÔ HÌNH, KIẾN TRÚC BẢO MẬT": "TRUY XUẤT ẢNH KẾT HỢP ZERO-SHOT CHỈ HUẤN LUYỆN BẰNG VĂN BẢN:",
    "CỦA HỆ ĐIỀU HÀNH WINDOWS SERVER 2025": "TÁI HIỆN VÀ CẢI TIẾN PHƯƠNG PHÁP LINCIR",
}
TEACHER = "TS. Đào Thị Thuý Quỳnh"
CLASS = "M25CQHT02-B"
GROUP = "8"
MEMBERS = [("B25CHHT112", "Đào Văn Tâm"), ("B25CHHT117", "Lâm Thành Trung"), ("B25CHHT119", "Nguyễn Xuân Tùng"),
           ("B25CHHT088", "Trần Quang Đức Dũng"), ("B25CHHT124", "Mekdala Nounou")]

# ---------------------------------------------------------------- docx helpers


def _texts(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def set_par_text(p_el, text):
    """Put `text` in the first run of a w:p (keeps its formatting), empty the others."""
    ts = list(p_el.iter(qn("w:t")))
    if not ts:
        return
    ts[0].text = text
    ts[0].set(qn("xml:space"), "preserve")
    for t in ts[1:]:
        t.text = ""


def edit_cover(doc):
    body = doc.element.body
    for p in body.iter(qn("w:p")):  # includes text boxes and their VML fallback copies
        txt = _texts(p).strip()
        for old, new in COVER.items():
            if txt == old:
                set_par_text(p, new)
    t_info, t_members = doc.tables[0], doc.tables[1]
    vals = {"Giảng viên:": TEACHER, "Lớp:": CLASS, "Nhóm:": GROUP}
    for row in t_info.rows:
        key = row.cells[0].text.strip()
        if key in vals:
            cell = row.cells[1]
            set_par_text(cell.paragraphs[0]._p, vals[key])
            for extra in cell.paragraphs[1:]:
                set_par_text(extra._p, "")
    # members: 2 per row, add rows as needed
    while len(t_members.rows) * 2 < len(MEMBERS):
        new_tr = copy.deepcopy(t_members.rows[-1]._tr)
        for bm in list(new_tr.iter(qn("w:bookmarkStart"))) + list(new_tr.iter(qn("w:bookmarkEnd"))):
            bm.getparent().remove(bm)  # bookmark ids must stay unique
        t_members._tbl.append(new_tr)
    cells = [c for r in t_members.rows for c in r.cells]
    for i, cell in enumerate(cells):
        paras = cell.paragraphs
        code, name = MEMBERS[i] if i < len(MEMBERS) else ("", "")
        if len(paras) >= 2:
            set_par_text(paras[0]._p, code)
            set_par_text(paras[1]._p, name)
            for extra in paras[2:]:
                set_par_text(extra._p, "")
        else:  # one paragraph "code<w:br/>name": code in the runs before the line break, name after it
            before, after, seen_br = [], [], False
            for el in paras[0]._p.iter():
                if el.tag == qn("w:br"):
                    seen_br = True
                elif el.tag == qn("w:t"):
                    (after if seen_br else before).append(el)
            for group, text in ((before, code), (after, name)):
                for j, t in enumerate(group):
                    t.text = text if j == 0 else ""
                    t.set(qn("xml:space"), "preserve")


def tighten_cover(doc):
    """Make room for the third row of members: drop the empty spacer paragraph between the title and the info table
    and the one after the members table. The spacers above the title stay, the PTIT logo is anchored there."""
    body = list(doc.element.body)
    tables = [i for i, el in enumerate(body) if el.tag == qn("w:tbl")][:2]
    for i in (tables[0] - 1, tables[1] + 1):
        el = body[i]
        if el.tag == qn("w:p") and not _texts(el).strip() and not list(el.iter(qn("w:drawing"))):
            el.getparent().remove(el)


def clear_after_cover(doc):
    body = doc.element.body
    start = None
    for i, el in enumerate(list(body)):
        if el.tag == qn("w:p") and _texts(el).strip() == "MỤC LỤC":
            start = i
            break
    assert start is not None, "template changed: 'MỤC LỤC' not found"
    for el in list(body)[start:]:
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def add_field(paragraph, instr, placeholder=""):
    def fld(kind):
        r = OxmlElement("w:r")
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), kind)
        r.append(f)
        return r
    paragraph._p.append(fld("begin"))
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r.append(it)
    paragraph._p.append(r)
    paragraph._p.append(fld("separate"))
    if placeholder:
        paragraph.add_run(placeholder)
    paragraph._p.append(fld("end"))


class Writer:
    def __init__(self, doc):
        self.doc = doc
        self.chapter = 0
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0

    # structure
    def title(self, text, new_page=True):
        p = self.doc.add_paragraph()
        p.paragraph_format.page_break_before = new_page
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(16)
        return p

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text, numbered=True):
        if numbered:
            self.chapter += 1
            self.fig_n = self.tab_n = self.eq_n = 0
            text = f"CHƯƠNG {self.chapter}. {text}"
        self.doc.add_paragraph(text, style="Heading 1").paragraph_format.page_break_before = True

    def h2(self, text):
        self.doc.add_paragraph(text, style="Heading 2")

    def h3(self, text):
        self.doc.add_paragraph(text, style="Heading 3")

    # text
    def p(self, text, indent=True):
        p = self.doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Cm(1)
        self._rich(p, text)
        return p

    def _rich(self, p, text):
        """**bold**, *italic*, [[placeholder]] (yellow)"""
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|\[\[[^\]]+\]\])", text):
            if not part:
                continue
            if part.startswith("**"):
                p.add_run(part[2:-2]).bold = True
            elif part.startswith("[["):
                r = p.add_run(f"[{part[2:-2]}]")
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif part.startswith("*"):
                p.add_run(part[1:-1]).italic = True
            else:
                p.add_run(part)

    def bullets(self, items):
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self._rich(p, it)

    def eq(self, text):
        self.eq_n += 1
        p = self.doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.italic = True
        p.add_run(f"        ({self.chapter}.{self.eq_n})")
        return f"({self.chapter}.{self.eq_n})"

    def placeholder(self, text):
        p = self.doc.add_paragraph(style="Normal")
        r = p.add_run(f"[CHỜ KẾT QUẢ] {text}")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        r.italic = True

    # floats
    def _caption(self, kind, text):
        k = self.fig_n if kind == "Hình" else self.tab_n
        p = self.doc.add_paragraph(style="Caption")
        p.add_run(f"{kind} {self.chapter}.").bold = True
        add_field(p, f"SEQ {kind} \\* ARABIC \\s 1", str(k))  # same scheme as the template; prefilled result
        for r in p.runs:
            if r.text == str(k):
                r.bold = True
        p.add_run(": ").bold = True
        p.add_run(text)
        return f"{kind} {self.chapter}.{k}"

    def figure(self, path, caption, width_cm=15.0):
        self.fig_n += 1
        p = self.doc.add_paragraph(style="Body Text")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if Path(path).exists():
            p.add_run().add_picture(str(path), width=Cm(width_cm))
        else:
            r = p.add_run(f"[CHỜ KẾT QUẢ] hình {Path(path).name} chưa có")
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return self._caption("Hình", caption)

    def table(self, header, rows, caption, widths_cm=None, font_pt=11):
        self.tab_n += 1
        label = self._caption("Bảng", caption)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(str(h))
            r.bold = True
            r.font.size = Pt(font_pt)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "D9E2F3")
            c._tc.get_or_add_tcPr().append(shd)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                s = vn(str(v))
                if s.startswith("[[") and s.endswith("]]"):
                    r = para.add_run(s[2:-2])
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                else:
                    r = para.add_run(s)
                r.font.size = Pt(font_pt)
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if (i == 0 or len(s) > 30) else WD_ALIGN_PARAGRAPH.CENTER
        for row in t.rows:  # keep body paragraphs compact inside tables
            for c in row.cells:
                for para in c.paragraphs:
                    para.paragraph_format.first_line_indent = Cm(0)
                    para.paragraph_format.space_after = Pt(0)
                    para.paragraph_format.line_spacing = 1.1
        if widths_cm:
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph()
        return label


def vn(text):
    """Vietnamese decimal comma: 26.35 -> 26,35 (only between digits)."""
    return re.sub(r"(\d)\.(\d)", r"\1,\2", text)


# ---------------------------------------------------------------- results (filled from results/assets.json)


def load_assets():
    return json.loads(ASSETS.read_text()) if ASSETS.exists() else None


def mean_std(v):
    v = [x for x in (v or []) if x is not None]
    if not v:
        return "[[chờ]]"
    return f"{v[0]:.2f}" if len(v) == 1 else f"{statistics.mean(v):.2f} ± {statistics.stdev(v):.2f}"


def m(v):
    v = [x for x in (v or []) if x is not None]
    return statistics.mean(v) if v else None


# ---------------------------------------------------------------- content


def front_matter(w):
    w.title("MỤC LỤC")
    add_field(w.doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "[Mở trong Word: chọn toàn bộ (Cmd/Ctrl+A) rồi nhấn F9 để cập nhật mục lục]")
    w.title("DANH MỤC HÌNH VẼ")
    add_field(w.doc.add_paragraph(), 'TOC \\h \\z \\c "Hình"', "[Cập nhật trường để hiện danh mục hình]")
    w.doc.add_paragraph()
    w.title("DANH MỤC BẢNG", new_page=False)
    add_field(w.doc.add_paragraph(), 'TOC \\h \\z \\c "Bảng"', "[Cập nhật trường để hiện danh mục bảng]")
    w.title("DANH MỤC TỪ VIẾT TẮT")
    abbr = [
        ("AP / mAP@K", "(mean) Average Precision at K", "Độ chính xác trung bình (trung bình trên các truy vấn) tại K"),
        ("CIR", "Composed Image Retrieval", "Truy xuất ảnh kết hợp (ảnh + văn bản)"),
        ("CIRCO", "CIR on Common Objects in context", "Bộ đánh giá CIR có nhiều đáp án đúng"),
        ("CIRR", "Composed Image Retrieval on Real-life images", "Bộ đánh giá CIR trên ảnh đời thực"),
        ("CLIP", "Contrastive Language–Image Pre-training", "Mô hình tiền huấn luyện đối sánh ảnh–văn bản"),
        ("FIQ", "FashionIQ", "Bộ đánh giá CIR trong lĩnh vực thời trang"),
        ("GPU", "Graphics Processing Unit", "Bộ xử lý đồ họa"),
        ("IDF", "Inverse Document Frequency", "Tần suất tài liệu nghịch đảo"),
        ("InfoNCE", "Info Noise-Contrastive Estimation", "Hàm mất mát đối sánh dựa trên phân loại mẫu dương/âm"),
        ("IR", "Information Retrieval", "Truy xuất thông tin"),
        ("KTC", "—", "Khoảng tin cậy"),
        ("LinCIR", "Language-only training for CIR", "Phương pháp CIR chỉ huấn luyện bằng văn bản"),
        ("MLP", "Multi-Layer Perceptron", "Mạng nơ-ron truyền thẳng nhiều lớp"),
        ("MSE", "Mean Squared Error", "Sai số bình phương trung bình"),
        ("R@K", "Recall at K", "Độ triệu hồi tại K"),
        ("SMP", "Self-Masking Projection", "Phép chiếu tự che (cơ chế tự giám sát của LinCIR)"),
        ("TB", "—", "Trung bình"),
        ("ViT", "Vision Transformer", "Kiến trúc Transformer cho ảnh"),
        ("ZS-CIR", "Zero-Shot CIR", "CIR không dùng dữ liệu bộ ba có nhãn"),
    ]
    w.tab_n = 0
    t = w.doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(("Viết tắt", "Tiếng Anh", "Ý nghĩa")):
        t.rows[0].cells[i].text = h
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
    for a in abbr:
        cells = t.add_row().cells
        for i, v in enumerate(a):
            cells[i].text = v
    for row in t.rows:
        for i, wd in enumerate((3.0, 6.0, 6.0)):
            row.cells[i].width = Cm(wd)
            for para in row.cells[i].paragraphs:
                para.paragraph_format.first_line_indent = Cm(0)
                para.paragraph_format.line_spacing = 1.0
                para.paragraph_format.space_after = Pt(0)
                for r in para.runs:
                    r.font.size = Pt(11)


def chapter1(w):
    w.h1("GIỚI THIỆU")
    w.h2("1.1. Bối cảnh và lý do chọn đề tài")
    w.p("Truy xuất thông tin (Information Retrieval – IR) là việc tìm các tài liệu, thường là phi cấu trúc, thỏa mãn một nhu cầu thông tin từ các bộ sưu tập lớn [12, 13]. Trong mô hình tìm kiếm cổ điển, nhu cầu thông tin của người dùng được diễn đạt thành câu truy vấn, hệ thống so khớp câu truy vấn với tập tài liệu đã được đánh chỉ mục và trả về danh sách kết quả; người dùng có thể tinh chỉnh truy vấn nếu kết quả chưa đạt [13]. Khi “tài liệu” là ảnh, câu truy vấn có thể là từ khóa, là một ảnh mẫu, hoặc là sự kết hợp của cả hai.")
    w.p("Nhiều nhu cầu thông tin thực tế khó diễn đạt chỉ bằng văn bản hoặc chỉ bằng ảnh. Một người mua sắm trực tuyến có thể muốn “chiếc váy giống ảnh này nhưng màu bạc, tay ngắn hơn”: ảnh mẫu mô tả phần lớn đặc điểm, còn câu văn chỉ mô tả phần cần thay đổi. Dạng truy vấn này được gọi là **truy xuất ảnh kết hợp** (Composed Image Retrieval – CIR). CIR mang lại mức độ tự do cao hơn cho người dùng và có nhiều ứng dụng trong thương mại điện tử, tìm kiếm ảnh trên web và các hệ thống hỏi đáp đa phương thức [1, 7, 8].")
    w.p("Đề tài được chọn vì ba lý do. Thứ nhất, CIR là một bài toán IR hiện đại, mở rộng trực tiếp các khái niệm của môn học (nhu cầu thông tin, truy vấn, xếp hạng, đánh giá Precision/Recall) sang dữ liệu đa phương thức. Thứ hai, bài báo LinCIR [1] (CVPR 2024) đề xuất một ý tưởng đơn giản nhưng hiệu quả: huấn luyện hệ thống truy xuất ảnh mà **chỉ cần dữ liệu văn bản**, có mã nguồn công khai, thuận lợi cho việc tái hiện. Thứ ba, hàm mất mát và cách thêm nhiễu của LinCIR còn những điểm có thể cải tiến và kiểm chứng được trong giới hạn tài nguyên của một nhóm học viên.")
    w.h2("1.2. Bài toán truy xuất ảnh kết hợp")
    w.p("Trong CIR, mỗi truy vấn gồm một **ảnh tham chiếu** x_ref và một **văn bản điều chỉnh** x_c mô tả sự khác biệt mong muốn; hệ thống cần trả về các ảnh trong kho thỏa mãn đồng thời cả hai điều kiện (Hình 1.1). Một bộ dữ liệu CIR có giám sát gồm các bộ ba ⟨ảnh tham chiếu, văn bản điều chỉnh, ảnh đích⟩.")
    w.figure(FIG / "cir_example.png", "Một truy vấn CIR thực tế trong FashionIQ: ảnh tham chiếu và văn bản điều chỉnh, ảnh đích cần tìm")
    w.p("Khó khăn lớn nhất của CIR là dữ liệu: các bộ ba không thể thu thập tự động bằng cách thu thập dữ liệu web mà phải gán nhãn thủ công, nên rất tốn kém và có quy mô nhỏ (FashionIQ khoảng 46,6 nghìn, CIRR khoảng 28,8 nghìn bộ ba huấn luyện) [1]. Các mô hình học có giám sát trên lượng dữ liệu này khó tổng quát hóa sang lĩnh vực mới. Vì vậy, hướng **CIR không cần mẫu** (zero-shot CIR – ZS-CIR) loại bỏ hoàn toàn sự phụ thuộc vào bộ ba: mô hình chỉ được huấn luyện trên dữ liệu sẵn có (ảnh hoặc văn bản) và được đánh giá trực tiếp trên các bộ dữ liệu CIR [2, 3].")
    w.h2("1.3. Mục tiêu và phạm vi")
    w.p("Mục tiêu của đề tài gồm bốn phần:")
    w.bullets([
        "**Đọc hiểu** bài báo LinCIR [1] và các phương pháp liên quan (Pic2Word [3], SEARLE [2]).",
        "**Dựng mã nguồn và tái hiện**: chạy được mã nguồn chính thức [21] trên môi trường hiện tại, đánh giá checkpoint chính thức và tự huấn luyện lại LinCIR.",
        "**Cải tiến**: đề xuất, cài đặt và huấn luyện các biến thể cải tiến; so sánh công bằng với baseline do nhóm tự chạy.",
        "**Chứng minh**: dùng kiểm định thống kê theo từng truy vấn, ablation và phân tích cơ chế (modality gap) để giải thích vì sao cải tiến tốt hơn hoặc không.",
    ])
    w.p("Do giới hạn tài nguyên (GPU NVIDIA T4 miễn phí trên Kaggle, khoảng 30 giờ mỗi tuần), đề tài sử dụng backbone CLIP ViT-L/14, huấn luyện 5.000 bước cho mọi biến thể (bài báo dùng 20.000 bước trên 8 GPU A100) và đánh giá trên FashionIQ, CIRCO. Bộ CIRR [7] không được dùng vì ảnh (thuộc NLVR2) cần đăng ký và quyền truy cập không kịp trong thời gian thực hiện đề tài.")
    w.h2("1.4. Đóng góp và cấu trúc báo cáo")
    w.p("Các đóng góp chính của nhóm gồm: (1) chuyển mã nguồn LinCIR sang môi trường hiện tại (Python 3.13, PyTorch 2.11, transformers 5.x) và kiểm chứng tính đúng đắn với sai lệch bằng 0 so với mô hình gốc; (2) hai cải tiến cho hàm mục tiêu và cơ chế nhiễu của LinCIR; (3) một quy trình thực nghiệm có kiểm định thống kê, có thể tái lập bằng một lệnh cho mỗi thí nghiệm.")
    w.p("Báo cáo gồm sáu chương. Chương 2 trình bày cơ sở lý thuyết và các nghiên cứu liên quan. Chương 3 phân tích phương pháp LinCIR và mã nguồn. Chương 4 trình bày các đề xuất cải tiến. Chương 5 mô tả thiết lập thực nghiệm, kết quả và phân tích. Chương 6 đánh giá ưu, nhược điểm và hạn chế; cuối cùng là kết luận và hướng phát triển.")


def chapter2(w):
    w.h1("CƠ SỞ LÝ THUYẾT VÀ NGHIÊN CỨU LIÊN QUAN")
    w.h2("2.1. Truy xuất thông tin và các độ đo đánh giá")
    w.p("Hiệu quả của một hệ thống IR được đánh giá bằng cách so sánh tập tài liệu hệ thống trả về với tập tài liệu thực sự liên quan đến nhu cầu thông tin [12, 13]. Hai độ đo cơ bản là độ chính xác (Precision) và độ triệu hồi (Recall):")
    w.eq("Precision = |Truy xuất ∩ Liên quan| / |Truy xuất|,    Recall = |Truy xuất ∩ Liên quan| / |Liên quan|")
    w.p("Các hệ thống CIR trả về một danh sách **có xếp hạng**, nên các độ đo được tính tại K vị trí đầu. Với các bộ dữ liệu chỉ có một ảnh đích cho mỗi truy vấn (FashionIQ, CIRR), độ đo thường dùng là **Recall@K**: tỷ lệ truy vấn có ảnh đích nằm trong K kết quả đầu tiên. Với CIRCO, mỗi truy vấn có nhiều ảnh đúng nên dùng **mAP@K** (mean Average Precision at K):")
    w.eq("AP@K = (1 / min(K, |G|)) · Σ_{k=1..K} P@k · rel(k),    mAP@K = trung bình AP@K trên mọi truy vấn")
    w.p("trong đó G là tập ảnh đúng, P@k là độ chính xác tại vị trí k và rel(k) bằng 1 nếu ảnh ở vị trí k thuộc G. mAP@K thưởng cho việc đặt nhiều ảnh đúng ở vị trí cao, do đó phản ánh chất lượng xếp hạng tốt hơn Recall@K. Bài báo LinCIR chỉ ra rằng R@1 trên các bộ một-đáp-án dễ bị nhiễu bởi “âm tính giả” (ảnh đúng nhưng không được gán nhãn) [1, 20], nên đề nghị ưu tiên mAP trên CIRCO và R@K với K lớn.")
    w.h2("2.2. Biểu diễn đa phương thức với CLIP")
    w.p("CLIP [4] gồm hai bộ mã hóa: bộ mã hóa ảnh I (Vision Transformer [25]) và bộ mã hóa văn bản C (Transformer [24]), được huấn luyện đối sánh trên 400 triệu cặp ảnh–văn bản để đưa ảnh và câu mô tả tương ứng về gần nhau trong một không gian nhúng chung d chiều. Văn bản được tách thành token, mỗi token được ánh xạ thành một **token embedding** e qua lớp nhúng E_w, rồi bộ mã hóa văn bản tạo **vector nhúng văn bản** z_c = C(E_w(t_c)); ảnh được mã hóa thành z_i = I(x_i). Độ tương đồng giữa văn bản và ảnh là cosine của hai vector.")
    w.p("Nhờ không gian chung này, truy xuất ảnh bằng văn bản trở thành bài toán **tìm kiếm vector**: mã hóa toàn bộ kho ảnh một lần, mã hóa truy vấn và xếp hạng theo cosine. Đây là một mô hình truy xuất có xếp hạng, khác với mô hình Boolean khớp chính xác dựa trên chỉ mục ngược [13].")
    w.h2("2.3. Khoảng cách giữa hai miền (modality gap)")
    w.p("Liang và cộng sự [5] chỉ ra rằng trong các mô hình như CLIP, vector nhúng ảnh và vector nhúng văn bản không trộn lẫn mà nằm ở hai vùng tách biệt (hai “hình nón”) của không gian nhúng. Khoảng cách này được đo bằng khoảng cách giữa tâm của các vector nhúng đã chuẩn hóa của hai miền. Hệ quả là một mô-đun chỉ được huấn luyện trên vector văn bản có thể hoạt động kém khi nhận vector ảnh ở thời điểm suy luận. Các phương pháp huấn luyện chỉ bằng văn bản cho bài toán sinh chú thích ảnh như CapDec [6] giảm vấn đề này bằng cách thêm nhiễu Gauss vào vector văn bản khi huấn luyện.")
    w.h2("2.4. Truy xuất ảnh kết hợp có giám sát")
    w.p("Các phương pháp CIR có giám sát học cách kết hợp đặc trưng ảnh tham chiếu và văn bản điều chỉnh từ các bộ ba có nhãn, ví dụ TIRG [10] hay Combiner [9] (kết hợp và tinh chỉnh một phần đặc trưng CLIP). Combiner là phương pháp có giám sát mạnh nhất trên FashionIQ được bài báo LinCIR so sánh. Hạn chế chung của hướng này là phụ thuộc vào dữ liệu bộ ba đắt đỏ và tổng quát hóa kém sang lĩnh vực khác.")
    w.h2("2.5. Zero-shot CIR dựa trên nghịch đảo văn bản")
    w.p("Hướng ZS-CIR chủ đạo dựa trên ý tưởng **nghịch đảo văn bản** (textual inversion [14]): biểu diễn ảnh tham chiếu thành một “từ giả” [$] trong không gian token embedding của CLIP, ghép vào câu “a photo of [$] that {văn bản điều chỉnh}”, rồi dùng bộ mã hóa văn bản tạo truy vấn và tìm ảnh như truy xuất văn bản → ảnh thông thường. Thành phần cần học là một mô-đun chiếu nhẹ φ ánh xạ vector ảnh z_i thành token embedding.")
    w.bullets([
        "**Pic2Word** [3] huấn luyện φ trên ảnh không nhãn bằng hàm mất mát đối sánh giữa z_i và vector nhúng của câu “a photo of [$]”.",
        "**SEARLE** [2] trước hết tối ưu một token giả cho từng ảnh (nghịch đảo dựa trên tối ưu), dùng GPT để sinh câu mô tả khái niệm rồi chưng cất thành mô-đun φ; SEARLE cũng giới thiệu bộ đánh giá CIRCO.",
    ])
    w.p("Theo phân tích của LinCIR [1], cả hai phương pháp có hai hạn chế: (1) chỉ dùng câu mẫu cố định “a photo of [$]” khi huấn luyện nên văn bản đầu vào thiếu đa dạng, hiệu năng giảm khi tăng kích thước backbone; (2) cần chạy bộ mã hóa ảnh khi huấn luyện, tốn tính toán và lưu trữ (ảnh CC3M khoảng 430 GB so với 125 MB chú thích).")


def chapter3(w):
    w.h1("PHƯƠNG PHÁP LINCIR")
    w.h2("3.1. Tổng quan")
    w.p("LinCIR (Language-only training for CIR) [1] giữ nguyên cách suy luận dựa trên token giả của Pic2Word/SEARLE nhưng thay đổi hoàn toàn cách huấn luyện: mô-đun chiếu φ được học **chỉ từ văn bản**. Thay vì chiếu vector ảnh z_i, LinCIR chiếu vector nhúng của chính câu văn z_c vào không gian token embedding và dùng một cơ chế tự giám sát gọi là **Self-Masking Projection** (SMP). Bộ mã hóa văn bản C được đóng băng; chỉ φ được huấn luyện. Hình 3.1 minh họa quá trình huấn luyện (các khối màu xanh lá là phần nhóm bổ sung ở Chương 4).")
    w.figure(FIG / "lincir_train.png", "Quá trình huấn luyện LinCIR (SMP) và vị trí hai cải tiến A, B của nhóm")
    w.h2("3.2. Self-Masking Projection")
    w.p("Cho một câu mô tả x_c (ví dụ “gray cat sleeps on a pillow”), LinCIR thực hiện các bước sau:")
    w.bullets([
        "Mã hóa câu: z_c = C(E_w(t_c)).",
        "Chiếu sang không gian token: ê_c = φ(z_c + n), với n là nhiễu ngẫu nhiên (mục 3.3).",
        "Xác định **từ khóa** là các cụm tính từ và danh từ liên tiếp bằng bộ gán nhãn từ loại của spaCy [22] (“gray cat”, “a pillow”), rồi thay toàn bộ từ khóa bằng [$]: “[$] sleeps on [$]”.",
        "Mã hóa câu đã che, trong đó token embedding của mọi vị trí [$] được thay bằng ê_c: ẑ_c = C(Ẽ).",
        "Tối thiểu hóa sai số bình phương giữa hai vector nhúng.",
    ])
    w.eq("L_SMP = ‖ z_c − ẑ_c ‖²")
    w.p("Trực giác của SMP: thông tin ngữ nghĩa của câu tập trung ở các từ khóa. Để câu đã che có cùng vector nhúng với câu gốc, token [$] phải “tóm tắt trong một từ” toàn bộ nội dung thiết yếu của câu. Do câu huấn luyện là văn bản tự nhiên đa dạng (chú thích CC3M, prompt Stable Diffusion), token [$] học cách hoạt động trong nhiều cấu trúc câu khác nhau, thay vì chỉ trong câu mẫu “a photo of [$]”. Ablation của bài báo cho thấy che toàn bộ từ khóa tốt hơn che token ngẫu nhiên, che các từ không phải từ khóa hay chỉ che danh từ [1, Bảng 8].")
    w.h2("3.3. Nhiễu để giảm khoảng cách giữa hai miền")
    w.p("Khi suy luận, φ nhận vector ảnh z_i chứ không phải z_c. Để φ tổng quát hóa qua modality gap, LinCIR thêm nhiễu vào z_c khi huấn luyện. Các tác giả nhận thấy nhiễu Gauss chuẩn trong không gian 768 chiều có chuẩn gần như cố định (hiện tượng “lời nguyền số chiều”), nên dùng một phân phối có chuẩn đa dạng hơn:")
    w.eq("n = u · ε,    u ~ U(0, 1),  ε ~ N(0, I)")
    w.p("Theo bài báo, lựa chọn này cho modality gap nhỏ nhất và hiệu năng tốt nhất trong các phân phối được thử (không nhiễu, Student-t, mũ, χ², N(0,1), U(−1,1)) [1, Bảng 1, 9].")
    w.h2("3.4. Suy luận")
    w.p("Khi truy vấn (Hình 3.2), ảnh tham chiếu được mã hóa thành z_i, chiếu thành token giả [$] = φ(z_i), ghép vào câu “a photo of [$] that {văn bản điều chỉnh}” và mã hóa bằng C để được vector truy vấn q. Kho ảnh được mã hóa trước bằng I; kết quả là danh sách ảnh xếp theo cosine(q, z_ảnh). Toàn bộ suy luận không cần huấn luyện thêm trên dữ liệu CIR.")
    w.figure(FIG / "lincir_infer.png", "Suy luận zero-shot CIR với mô-đun chiếu φ đã huấn luyện")
    w.h2("3.5. Hiện thực trong mã nguồn và các điểm khác biệt")
    w.p("Mã nguồn chính thức [21] hiện thực φ bằng MLP ba lớp tuyến tính (kích thước ẩn 4d, GELU, dropout 0,5), tối ưu bằng AdamW với tốc độ học 10⁻⁴, weight decay 0,01, batch 512 và 20.000 bước; checkpoint tốt nhất được chọn theo R@1 trên tập dev của CIRR. Khi đọc mã nguồn, nhóm ghi nhận ba điểm khác với mô tả trong bài báo:")
    w.bullets([
        "Bài báo mô tả φ có LayerNorm ở đầu và cuối (LN–Linear–…–Linear–LN), nhưng mã nguồn và checkpoint công bố **không có LayerNorm** (nhóm đã kiểm chứng: checkpoint ViT-L nạp khớp hoàn toàn vào lớp MLP không LayerNorm).",
        "Bài báo nêu dữ liệu huấn luyện gồm chú thích CC3M và 2,47 triệu prompt Stable Diffusion, nhưng mã nguồn nạp thêm một bộ **prompt Midjourney**.",
        "Hàm trích từ khóa trong mã nguồn giữ cả mạo từ/từ hạn định (DET) trong cụm từ khóa.",
    ])
    w.p("Nhóm giữ nguyên hành vi của mã nguồn chính thức cho baseline (vì đây là cấu hình sinh ra checkpoint công bố) và ghi nhận các điểm trên như nguồn chênh lệch tiềm năng khi so sánh với số liệu trong bài báo.")


def chapter4(w):
    w.h1("ĐỀ XUẤT CẢI TIẾN")
    w.h2("4.1. Phân tích hạn chế của LinCIR")
    w.p("Nhóm xác định hai điểm có thể cải thiện trong quá trình huấn luyện của LinCIR:")
    w.bullets([
        "**Hàm mục tiêu không phân biệt.** L_SMP chỉ kéo ẑ_c về gần z_c của chính câu đó. MSE không phạt việc ẑ_c cũng gần vector của các câu khác, trong khi bài toán đánh giá là **xếp hạng**: token [$] cần làm cho truy vấn gần ảnh đích hơn mọi ảnh khác. Ngoài ra, MSE trên vector chưa chuẩn hóa chịu ảnh hưởng của chuẩn vector, trong khi truy xuất chỉ dùng hướng (cosine).",
        "**Nhiễu đẳng hướng.** Nhiễu u·ε phân bố đều theo mọi hướng. Tuy nhiên, modality gap của CLIP có hướng rõ ràng: vector ảnh lệch khỏi vector văn bản theo một hướng gần như cố định [5]. Nhiễu đẳng hướng chỉ “làm mờ” vùng văn bản chứ không đưa đầu vào của φ về phía vùng ảnh mà φ sẽ gặp khi suy luận.",
    ])
    w.h2("4.2. Cải tiến A: bổ sung hàm mất mát đối sánh InfoNCE")
    w.p("Nhóm bổ sung một số hạng đối sánh InfoNCE [11] đối xứng trong batch. Với batch B câu, gọi ẑ và z là các vector đã chuẩn hóa L2:")
    w.eq("L_NCE = ½ [ CE(ẑ zᵀ / τ, I) + CE(z ẑᵀ / τ, I) ],    L = L_SMP + λ · L_NCE")
    w.p("trong đó CE là entropy chéo với nhãn là đường chéo (mỗi câu chỉ khớp với chính nó), τ = 0,07 là nhiệt độ và λ = 0,5. Với batch 512, mỗi câu có 511 mẫu âm. Số hạng này buộc token [$] mang thông tin **phân biệt** được câu này với các câu khác, đúng với bản chất xếp hạng của bài toán đánh giá, đồng thời tối ưu theo cosine như khi truy xuất. MSE vẫn được giữ để bảo toàn chuẩn và vị trí tuyệt đối của vector nhúng.")
    w.p("Chi phí tính toán gần như không đổi vì ẑ và z đã có sẵn; phép nhân ma trận 512 × 512 không đáng kể. Đo trên GPU T4, thời gian mỗi bước tăng khoảng 5% (2,59 so với 2,46 giây).")
    w.h2("4.3. Cải tiến B: nhiễu có hướng theo modality gap")
    w.p("Nhóm tính hướng của modality gap từ N cặp ảnh–chú thích (5.000 cặp COCO val2017 [17]) trong không gian mà φ nhận đầu vào (vector chưa chuẩn hóa):")
    w.eq("Δ = (1/N) Σ z_i − (1/N) Σ z_c,    n_B = α · Δ + u · ε")
    w.p("Nhiễu mới dịch tâm phân phối đầu vào của φ về phía vùng ảnh (α = 0,5), trong khi vẫn giữ thành phần ngẫu nhiên của LinCIR. Ý tưởng tương tự các phương pháp “dịch tâm” để thu hẹp modality gap [5], nhưng được áp dụng vào quá trình huấn luyện chỉ bằng văn bản. Δ chỉ cần tính một lần; dữ liệu ảnh dùng để tính Δ (COCO) độc lập với các bộ đánh giá.")
    w.h2("4.4. Hướng dự phòng C: chọn từ khóa theo IDF")
    w.p("Nhóm cũng cài đặt một biến thể chỉ che các từ khóa có giá trị IDF cao nhất (tỷ lệ r), dựa trên trọng số TF-IDF của môn học [12], nhằm tập trung [$] vào các khái niệm mang nhiều thông tin. Do giới hạn giờ GPU, hướng C chưa được đưa vào đợt thực nghiệm chính và được để lại cho hướng phát triển.")
    w.h2("4.5. Giả thuyết kiểm chứng")
    w.bullets([
        "**H1 (A):** thêm InfoNCE làm tăng các độ đo xếp hạng (FashionIQ R@10/R@50, CIRCO mAP@K) so với baseline cùng cấu hình.",
        "**H2 (B):** nhiễu có hướng làm φ cho kết quả nhất quán hơn giữa đầu vào ảnh và văn bản (cosine cao hơn, khoảng cách tâm nhỏ hơn) và từ đó tăng hiệu năng truy xuất.",
        "**H3 (A+B):** hai cải tiến tác động lên hai thành phần khác nhau (hàm mục tiêu và phân phối đầu vào) nên có thể cộng hưởng.",
    ])


def chapter5(w, A):
    w.h1("THỰC NGHIỆM VÀ ĐÁNH GIÁ")
    w.h2("5.1. Môi trường và dựng mã nguồn")
    w.p("Mã nguồn chính thức [21] được viết cho transformers 4.34 và phụ thuộc gói clip của OpenAI. Trên môi trường hiện tại (Kaggle: Python 3.13, PyTorch 2.11, transformers 5.16; máy cá nhân Mac mini M1: Python 3.14, transformers 5.19), nhóm phải xử lý các vấn đề sau:")
    w.bullets([
        "Hàm mã hóa câu có token giả gọi trực tiếp API nội bộ của transformers đã thay đổi: nhóm viết lại lượt truyền xuôi của bộ mã hóa văn bản từ các mô-đun con và **kiểm chứng sai lệch bằng 0** so với CLIPTextModel gốc trên cả ViT-B/32 (máy cá nhân) và ViT-L/14 (GPU Kaggle).",
        "Thay clip.tokenize bằng tokenizer CLIP của HuggingFace (cùng bộ từ vựng BPE, token “$” có cùng mã 259).",
        "Sửa các lỗi trong mã gốc: biến gõ nhầm khi cắt gradient, tham số sai khi đánh giá CIRCO, ép kiểu .half() gây lỗi trên CPU/MPS, ghi tham số không hợp lệ vào TensorBoard.",
        "Bổ sung công cụ: đánh giá lưu thứ hạng từng truy vấn, kiểm định thống kê, đo modality gap, chạy song song hai GPU và sổ đăng ký thí nghiệm.",
    ])
    w.p("Huấn luyện và đánh giá chạy trên Kaggle với 2 GPU NVIDIA T4 (16 GB). Mỗi biến thể chạy trên một GPU với batch 512, fp16; bộ nhớ đỉnh đo được là 11,4 GB, tốc độ khoảng 2,4 giây/bước.")
    w.h2("5.2. Dữ liệu")
    w.p("**Dữ liệu huấn luyện** là văn bản thuần, giống mã nguồn chính thức: chú thích CC3M [16], 2,47 triệu prompt Stable Diffusion và prompt Midjourney, tải tự động từ HuggingFace. **Dữ liệu đánh giá** gồm hai bộ không cần đăng ký (Bảng 5.1). Ảnh FashionIQ lấy từ bản sao trên HuggingFace (royokong/fashioniq_val_imgs) do phần lớn liên kết ảnh gốc đã hỏng; nhóm kiểm tra đủ 15.415/15.415 ảnh cần cho tập val. CIRCO dùng ảnh COCO 2017 unlabeled (123.403 ảnh) [17]; chỉ tập val có đáp án công khai.")
    w.table(["Bộ dữ liệu", "Lĩnh vực", "Số truy vấn", "Kho ảnh", "Số ảnh đúng / truy vấn", "Độ đo"],
            [["FashionIQ val – dress", "Thời trang", "2.017", "3.817", "1", "R@10, R@50"],
             ["FashionIQ val – shirt", "Thời trang", "2.038", "6.346", "1", "R@10, R@50"],
             ["FashionIQ val – toptee", "Thời trang", "1.961", "5.373", "1", "R@10, R@50"],
             ["CIRCO val", "Ảnh đời thực (COCO)", "220", "123.403", "4,16 (TB)", "mAP@5/10/25/50"]],
            "Các bộ dữ liệu đánh giá", widths_cm=[3.6, 2.6, 2.0, 2.0, 2.6, 2.6])
    w.h2("5.3. Giao thức thực nghiệm")
    w.bullets([
        "**Biến thể:** baseline (LinCIR như mã nguồn chính thức), A (+InfoNCE, λ = 0,5), B (+nhiễu có hướng, α = 0,5), A+B. Ngoài ra đánh giá checkpoint chính thức để đối chiếu với bài báo.",
        "**So sánh công bằng:** mọi biến thể dùng cùng backbone ViT-L/14, cùng dữ liệu, cùng 5.000 bước và cùng siêu tham số khác; chỉ khác đúng thành phần cải tiến.",
        "**Chọn checkpoint:** dùng checkpoint ở bước cuối cho mọi biến thể; không chọn theo FashionIQ hay CIRCO để tránh chọn trên chính tập báo cáo.",
        "**Lặp lại:** mỗi biến thể chạy 3 seed (12345, 2024, 7), báo cáo trung bình ± độ lệch chuẩn.",
        "**Kiểm định:** kết quả được ghép theo từng truy vấn và lấy trung bình qua seed; dùng paired bootstrap 10.000 lần [18] và kiểm định Wilcoxon có dấu [19]. Chỉ kết luận “tốt hơn” khi p < 0,05.",
        "**Đường học:** đánh giá thêm các checkpoint 1.000–4.000 bước của baseline và A (seed 12345) để kiểm tra 5.000 bước có đủ hội tụ.",
    ])
    results_sections(w, A)


def results_sections(w, A):
    main = (A or {}).get("main", {})
    w.h2("5.4. Kết quả tái hiện")
    off = main.get("EXP-000-official", {})
    paper = {"dress": (20.92, 42.44), "shirt": (29.10, 46.81), "toptee": (28.81, 50.18), "average": (26.28, 46.49)}
    w.p("Trước khi huấn luyện, nhóm đánh giá lại checkpoint LinCIR ViT-L chính thức [21] bằng mã nguồn đã chuyển đổi và so sánh với Bảng 5 của bài báo [1] (Bảng 5.2). Đây là phép kiểm tra toàn bộ đường ống: dữ liệu FashionIQ lấy từ bản sao, bộ mã hóa văn bản viết lại, tokenizer thay thế và script đánh giá mới.")
    keys = [("dress", "Dress"), ("shirt", "Shirt"), ("toptee", "Toptee"), ("average", "TB")]
    head = ["Nguồn"] + [f"{n} R@{k}" for _, n in keys for k in (10, 50)]
    ours = []
    for d, _ in keys:
        for i, k in enumerate((10, 50)):
            v = (off.get(f"{d}_fiq_recall_at{k}") or [None])[0] if off else None
            ours.append(f"{v:.2f}" if v is not None else "[[chờ]]")
    diffs = []
    if off:
        diffs.append(["Chênh lệch"] + [f"{float(ours[2 * j + i]) - paper[d][i]:+.2f}" for j, (d, _) in enumerate(keys) for i in range(2)])
    w.table(head, [["Bài báo [1]"] + [f"{paper[d][i]:.2f}" for d, _ in keys for i in range(2)], ["Nhóm đánh giá lại"] + ours] + diffs,
            "Đối chiếu checkpoint LinCIR ViT-L chính thức với bài báo trên FashionIQ val (%)", widths_cm=[3.2] + [1.475] * 8, font_pt=9)
    if off:
        dif = [abs(float(x)) for x in diffs[0][1:]]
        c5 = (off.get("circo_map_at5") or [None])[0]
        c10 = (off.get("circo_map_at10") or [None])[0]
        w.p(vn(f"Kết quả tái hiện khớp rất sát bài báo: chênh lệch tuyệt đối lớn nhất là {max(dif):.2f} điểm phần trăm, trung bình FashionIQ R@10 và R@50 lệch {float(diffs[0][7]):+.2f} và {float(diffs[0][8]):+.2f} điểm. Mức lệch nhỏ này nằm trong phạm vi có thể giải thích bởi khác biệt phần mềm (phiên bản thư viện, tokenizer, độ chính xác số học) và ảnh FashionIQ từ bản sao. Điều đó xác nhận mã nguồn sau khi chuyển đổi và quy trình đánh giá của nhóm là đúng. Trên CIRCO val, checkpoint chính thức đạt mAP@5 = {c5:.2f} và mAP@10 = {c10:.2f}; bài báo chỉ công bố số trên tập test (12,59 và 13,58) nên hai con số không so sánh trực tiếp được, và số CIRCO val này được dùng làm mốc tham chiếu cho các biến thể của nhóm."))
    else:
        w.placeholder("Nhận xét mức độ khớp với bài báo sau khi có kết quả EXP-000 (phiên prep_eval).")
    w.h2("5.5. Kết quả chính")
    rows = []
    for prefix, label in (("EXP-001-baseline", "Baseline (LinCIR, 5k bước)"), ("EXP-002-infonce", "A: + InfoNCE"),
                          ("EXP-003-gapnoise", "B: + nhiễu có hướng"), ("EXP-004-infonce-gapnoise", "A + B")):
        r = main.get(prefix, {})
        rows.append([label] + [mean_std(r.get(k)) if r else "[[chờ]]" for k in
                               ("dress_fiq_recall_at10", "shirt_fiq_recall_at10", "toptee_fiq_recall_at10",
                                "average_fiq_recall_at10", "average_fiq_recall_at50", "circo_map_at5", "circo_map_at10")])
    w.table(["Biến thể", "Dress R@10", "Shirt R@10", "Toptee R@10", "FIQ TB R@10", "FIQ TB R@50", "CIRCO mAP@5", "CIRCO mAP@10"],
            rows, "Kết quả chính (trung bình ± độ lệch chuẩn qua 3 seed, ViT-L/14, 5.000 bước)",
            widths_cm=[3.4, 1.7, 1.7, 1.7, 1.7, 1.7, 1.7, 1.7], font_pt=9)
    if not main.get("EXP-001-baseline"):
        w.placeholder("Nhận xét bảng kết quả chính: biến thể nào tốt nhất, mức chênh so với baseline trên từng bộ dữ liệu.")
    w.h2("5.6. Kiểm định ý nghĩa thống kê và ablation")
    sig = (A or {}).get("significance", {})
    srows = []
    for prefix, label in (("EXP-002-infonce", "A"), ("EXP-003-gapnoise", "B"), ("EXP-004-infonce-gapnoise", "A+B")):
        for r in sig.get(prefix, []):
            if r["metric"] in ("R@10", "R@50", "ap5", "ap10"):
                pw = "—" if r["p_wilcoxon"] is None else f"{r['p_wilcoxon']:.4f}"
                srows.append([label, r["dataset"], r["metric"].replace("ap", "mAP@"), f"{r['delta']:+.2f}",
                              f"[{r['ci95'][0]:+.2f}; {r['ci95'][1]:+.2f}]", f"{r['p_boot']:.4f}", pw])
    if not srows:
        srows = [[x, "FashionIQ / CIRCO", "R@10, R@50, mAP@5, mAP@10", "[[chờ]]", "[[chờ]]", "[[chờ]]", "[[chờ]]"] for x in ("A", "B", "A+B")]
    w.table(["So với baseline", "Bộ dữ liệu", "Chỉ số", "Δ", "KTC 95%", "p (bootstrap)", "p (Wilcoxon)"], srows,
            "Kiểm định theo từng truy vấn (trung bình qua seed)", widths_cm=[2.2, 2.3, 2.6, 1.5, 3.0, 2.0, 2.0], font_pt=10)
    if not sig:
        w.placeholder("Kết luận H1–H3 dựa trên p < 0,05; nêu rõ chỉ số nào có/không có ý nghĩa thống kê.")
    w.p("Bảng 5.3 đồng thời là một ablation, tách đóng góp riêng của từng thành phần: so sánh A với baseline cho thấy tác dụng của InfoNCE, B với baseline cho tác dụng của nhiễu có hướng, A+B với A và B cho thấy hai thành phần có cộng hưởng hay không.")
    w.h2("5.7. Phân tích nguyên nhân")
    w.figure(RES_FIG / "modality_gap.png", "Độ nhất quán của φ giữa đầu vào ảnh và văn bản trên 5.000 cặp COCO val2017 (seed 12345)")
    if not (A or {}).get("modality_gap"):
        w.placeholder("Phân tích: biến thể nào làm φ(ảnh) và φ(văn bản) gần nhau hơn; có tương quan với mức tăng hiệu năng (H2) không.")
    w.figure(RES_FIG / "learning_curve.png", "Đường học của baseline và A trên FashionIQ (seed 12345)")
    if not (A or {}).get("curve"):
        w.placeholder("Nhận xét đường học: hiệu năng đã bão hòa ở 5.000 bước chưa; khoảng cách A – baseline có ổn định theo số bước không.")
    w.h2("5.8. Ví dụ định tính")
    q = (A or {}).get("qualitative") or []
    if q:
        for item in q[:2]:
            w.figure(RES_FIG / item["file"], f"Top-5 kết quả của baseline và biến thể cải tiến cho một truy vấn FashionIQ ({item['dress']})")
    else:
        w.figure(RES_FIG / "qualitative_fiq_1.png", "Top-5 kết quả của baseline và biến thể cải tiến cho một truy vấn FashionIQ")
        w.placeholder("Bình luận ví dụ thành công và ví dụ thất bại.")


def chapter6(w):
    w.h1("ĐÁNH GIÁ VÀ HẠN CHẾ")
    w.h2("6.1. Ưu điểm")
    w.bullets([
        "LinCIR chỉ cần văn bản để huấn luyện: rẻ về lưu trữ, nhanh, dễ mở rộng sang backbone lớn; nhóm tái hiện được trên GPU miễn phí.",
        "Hai cải tiến đề xuất không thay đổi kiến trúc hay cách suy luận, chi phí huấn luyện tăng không đáng kể và bật/tắt được bằng một tham số.",
        "Quy trình thực nghiệm có kiểm định thống kê theo truy vấn, chạy lại được bằng một lệnh cho mỗi thí nghiệm, mã nguồn được kiểm chứng tương đương mô hình gốc.",
    ])
    w.h2("6.2. Hạn chế")
    w.bullets([
        "Huấn luyện 5.000 bước thay vì 20.000 bước như bài báo do giới hạn giờ GPU; số tuyệt đối của baseline có thể thấp hơn checkpoint chính thức.",
        "Chưa quét siêu tham số λ, α; các giá trị 0,5 có thể chưa tối ưu.",
        "Không đánh giá trên CIRR (ảnh NLVR2 cần đăng ký, quyền truy cập không kịp thời hạn), nên không đối chiếu được Bảng 6 của bài báo; bài báo cũng nhận xét CIRR có nhiều nhiễu (âm tính giả, câu điều chỉnh không thật sự mang tính “điều chỉnh”) [1]. CIRCO chỉ dùng tập val (220 truy vấn) do tập test chấm trên máy chủ.",
        "Ảnh FashionIQ lấy từ bản sao cộng đồng; dữ liệu huấn luyện chứa prompt Midjourney như mã nguồn chính thức nhưng khác mô tả trong bài báo.",
    ])
    w.h1("KẾT LUẬN", numbered=False)
    w.p("Đề tài đã nghiên cứu bài toán truy xuất ảnh kết hợp dưới góc nhìn của truy xuất thông tin và phương pháp LinCIR huấn luyện chỉ bằng văn bản. Nhóm đã chuyển mã nguồn chính thức sang môi trường hiện tại với sai lệch bằng 0 so với mô hình gốc, dựng quy trình huấn luyện – đánh giá – kiểm định có thể tái lập, và đề xuất hai cải tiến: bổ sung hàm mất mát đối sánh InfoNCE vào cơ chế Self-Masking Projection, và dùng nhiễu có hướng theo modality gap.")
    w.placeholder("Tóm tắt kết quả chính và kết luận cho H1–H3 (điền sau khi có số liệu).")
    w.p("Hướng phát triển gồm: huấn luyện đủ 20.000 bước và quét λ, α; đánh giá trên CIRR và tập test của CIRCO khi có quyền truy cập; thử hướng C (chọn từ khóa theo IDF) và mở rộng sang backbone lớn hơn.")


REFERENCES = [
    "G. Gu, S. Chun, W. Kim, Y. Kang, S. Yun, “Language-only Efficient Training of Zero-shot Composed Image Retrieval,” in Proc. CVPR, 2024, pp. 13225–13234.",
    "A. Baldrati, L. Agnolucci, M. Bertini, A. Del Bimbo, “Zero-Shot Composed Image Retrieval with Textual Inversion,” in Proc. ICCV, 2023.",
    "K. Saito, K. Sohn, X. Zhang, C.-L. Li, C.-Y. Lee, K. Saenko, T. Pfister, “Pic2Word: Mapping Pictures to Words for Zero-shot Composed Image Retrieval,” in Proc. CVPR, 2023.",
    "A. Radford et al., “Learning Transferable Visual Models From Natural Language Supervision,” in Proc. ICML, 2021.",
    "W. Liang, Y. Zhang, Y. Kwon, S. Yeung, J. Zou, “Mind the Gap: Understanding the Modality Gap in Multi-modal Contrastive Representation Learning,” in Proc. NeurIPS, 2022.",
    "D. Nukrai, R. Mokady, A. Globerson, “Text-Only Training for Image Captioning using Noise-Injected CLIP,” in Findings of EMNLP, 2022.",
    "Z. Liu, C. Rodriguez-Opazo, D. Teney, S. Gould, “Image Retrieval on Real-life Images with Pre-trained Vision-and-Language Models,” in Proc. ICCV, 2021.",
    "H. Wu et al., “Fashion IQ: A New Dataset Towards Retrieving Images by Natural Language Feedback,” in Proc. CVPR, 2021.",
    "A. Baldrati, M. Bertini, T. Uricchio, A. Del Bimbo, “Conditioned and Composed Image Retrieval Combining and Partially Fine-Tuning CLIP-Based Features,” in Proc. CVPR Workshops, 2022.",
    "N. Vo et al., “Composing Text and Image for Image Retrieval – An Empirical Odyssey,” in Proc. CVPR, 2019.",
    "A. van den Oord, Y. Li, O. Vinyals, “Representation Learning with Contrastive Predictive Coding,” arXiv:1807.03748, 2018.",
    "C. D. Manning, P. Raghavan, H. Schütze, Introduction to Information Retrieval, Cambridge University Press, 2008.",
    "Đào Thị Thuý Quỳnh, Bài giảng môn Tìm kiếm và Truy xuất thông tin – Chương I: Các khái niệm cơ bản, Học viện Công nghệ Bưu chính Viễn thông.",
    "R. Gal et al., “An Image is Worth One Word: Personalizing Text-to-Image Generation using Textual Inversion,” in Proc. ICLR, 2023.",
    "G. Ilharco et al., OpenCLIP, Zenodo, 2021. https://github.com/mlfoundations/open_clip",
    "P. Sharma, N. Ding, S. Goodman, R. Soricut, “Conceptual Captions: A Cleaned, Hypernymed, Image Alt-text Dataset for Automatic Image Captioning,” in Proc. ACL, 2018.",
    "T.-Y. Lin et al., “Microsoft COCO: Common Objects in Context,” in Proc. ECCV, 2014.",
    "B. Efron, R. J. Tibshirani, An Introduction to the Bootstrap, Chapman & Hall/CRC, 1993.",
    "F. Wilcoxon, “Individual Comparisons by Ranking Methods,” Biometrics Bulletin, vol. 1, no. 6, pp. 80–83, 1945.",
    "S. Chun, W. Kim, S. Park, M. Chang, S. J. Oh, “ECCV Caption: Correcting False Negatives by Collecting Machine-and-Human-verified Image-Caption Associations for MS-COCO,” in Proc. ECCV, 2022.",
    "NAVER Vision, “LinCIR: Official PyTorch implementation,” GitHub. https://github.com/navervision/lincir",
    "M. Honnibal, I. Montani, S. Van Landeghem, A. Boyd, spaCy: Industrial-strength Natural Language Processing in Python, 2020.",
    "T. Wolf et al., “Transformers: State-of-the-Art Natural Language Processing,” in Proc. EMNLP (System Demonstrations), 2020.",
    "A. Vaswani et al., “Attention Is All You Need,” in Proc. NeurIPS, 2017.",
    "A. Dosovitskiy et al., “An Image is Worth 16x16 Words: Transformers for Image Recognition at Scale,” in Proc. ICLR, 2021.",
]


def references(w):
    w.h1("TÀI LIỆU THAM KHẢO", numbered=False)
    for i, ref in enumerate(REFERENCES, start=1):
        p = w.doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Cm(1)
        p.paragraph_format.first_line_indent = Cm(-1)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"[{i}]\t{ref}").font.size = Pt(12)


def appendix(w):
    w.h1("PHỤ LỤC. PHÂN CÔNG CÔNG VIỆC", numbered=False)
    w.chapter = "PL"
    w.tab_n = 0
    w.table(["Thành viên", "Nhiệm vụ chính"],
            [["Đào Văn Tâm (B25CHHT112)", "Chuyển đổi mã nguồn LinCIR sang môi trường mới, kiểm chứng tương đương với mô hình gốc; vận hành thực nghiệm trên Kaggle: lập lịch chạy song song trong hạn mức GPU, khắc phục nghẽn nạp dữ liệu."],
             ["Lâm Thành Trung (B25CHHT117)", "Đề xuất và cài đặt cải tiến A (bổ sung InfoNCE vào SMP); phân tích đường học theo số bước huấn luyện."],
             ["Nguyễn Xuân Tùng (B25CHHT119)", "Tái hiện: đánh giá checkpoint chính thức, đối chiếu với bài báo; phân tích khác biệt giữa bài báo và mã nguồn (cấu trúc φ, dữ liệu huấn luyện, trích từ khóa)."],
             ["Trần Quang Đức Dũng (B25CHHT088)", "Thiết kế giao thức đánh giá theo truy vấn (R@K, mAP@K) và kiểm định thống kê (paired bootstrap, Wilcoxon); cài đặt cải tiến B (nhiễu có hướng) và phân tích modality gap."],
             ["Mekdala Nounou (B25CHHT124)", "Tổng quan lý thuyết (CIR, ZS-CIR, CLIP, Pic2Word, SEARLE) và liên hệ với các độ đo IR; phân tích định tính; thiết kế slides thuyết trình."]],
            "Phân công công việc trong nhóm", widths_cm=[4.3, 10.7], font_pt=11)


STAMP = OUT.with_suffix(".generated.sha256")  # hash of the last generated file: detects manual edits


def _sha(path):
    import hashlib
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    global OUT
    if OUT.exists() and (not STAMP.exists() or STAMP.read_text().strip() != _sha(OUT)):
        # the .docx was edited by hand (Word) after the last build: never overwrite it
        OUT = OUT.with_name(OUT.stem + "_auto.docx")
        print(f"[!] Báo cáo đã được sửa tay -> không ghi đè; bản tự sinh ghi vào {OUT.name}. Chép phần cần thiết sang bản sửa tay.")
    shutil.copy(TEMPLATE, OUT)
    doc = docx.Document(OUT)
    edit_cover(doc)
    tighten_cover(doc)
    clear_after_cover(doc)
    w = Writer(doc)
    A = load_assets()
    front_matter(w)
    chapter1(w)
    chapter2(w)
    chapter3(w)
    chapter4(w)
    chapter5(w, A)
    chapter6(w)
    references(w)
    appendix(w)
    doc.save(OUT)  # fields (TOC, lists) are refreshed by report/update_fields.sh (Word) or F9 in Word
    if OUT.name == "BaoCao_Nhom8_LinCIR.docx":
        STAMP.write_text(_sha(OUT))
    print(OUT, "results:", "có" if A else "chưa có (placeholder)")


if __name__ == "__main__":
    main()
