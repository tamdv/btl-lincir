"""[NHOM8] Restructure the hand-edited 15-slide deck for a single presenter (Tâm), following the lecturer-style review:
problem and research question first, metrics with formulas, token idea, paper summary, SMP (original only), noise and
modality gap, improvements A/B with hypotheses (no early results), code + Kaggle, reproduction + paper/code differences,
results with the official checkpoint row, analysis with a takeaway, conclusion with H1-H3 and a retrieval example.
Existing slides are edited in place (style, Morph names kept); one slide is added, one removed. Backup written first.

  python slides/restructure_solo.py [--deck slides/Slides_Nhom8_LinCIR.pptx]
"""
import argparse
import copy
import datetime as dt
import shutil
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent))
from add_overview_slides import frame  # noqa: E402
from build_slides import BLUE_DARK, NAVY, TINT, card, stat, text  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
B = {"bold": True, "color": BLUE_DARK}


def find(slide, name):
    return next(sh for sh in slide.shapes if sh.name == name)


def drop(slide, *names):
    for sh in list(slide.shapes):
        if sh.name in names:
            sh._element.getparent().remove(sh._element)


def set_paras(sh, paras):
    """replace paragraph texts in place, keeping each paragraph's and run's formatting"""
    tf = sh.text_frame
    tmpl = tf.paragraphs[-1]._p
    while len(tf.paragraphs) < len(paras):
        tmpl.addnext(copy.deepcopy(tmpl))
        tmpl = tf.paragraphs[-1]._p
    for p in list(tf.paragraphs)[len(paras):]:
        p._p.getparent().remove(p._p)
    for p, txt in zip(tf.paragraphs, paras):
        parts = txt if isinstance(txt, list) else [txt]
        runs = p.runs
        while len(runs) < len(parts):
            runs[-1]._r.addnext(copy.deepcopy(runs[-1]._r))
            runs = p.runs
        for r, t in zip(runs, parts):
            r.text = t
        for r in runs[len(parts):]:
            r._r.getparent().remove(r._r)


def set_title(slide, t):
    set_paras(find(slide, "!!title"), [t])


def fit_picture(slide, name, path, x, y, w, h):
    """replace a picture by another image, fitted (aspect kept) into the box x, y, w, h (inches)"""
    old = find(slide, name)
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih)
    pw, ph = iw * sc, ih * sc
    pic = slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y), Inches(pw), Inches(ph))
    old._element.addprevious(pic._element)
    old._element.getparent().remove(old._element)
    return pic


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", default=str(ROOT / "slides" / "Slides_Nhom8_LinCIR.pptx"))
    deck = Path(ap.parse_args().deck)
    prs = Presentation(deck)
    S = list(prs.slides)
    if len(S) != 15 or "Nhiễu giúp φ" in "".join(sh.text_frame.text for s in S for sh in s.shapes if sh.has_text_frame):
        sys.exit("Deck không ở trạng thái 15 slide trước khi tái cấu trúc; dừng để không sửa hai lần.")
    shutil.copy(deck, deck.with_name(f"{deck.stem}_truocKhiMotNguoi_{dt.datetime.now():%Y%m%d-%H%M}.pptx"))
    o = {i + 1: s for i, s in enumerate(S)}

    # N2 (old 3): problem + paper goal + research question
    s = o[3]
    text(s, 0.6, 4.95, 8.3, 1.75, [
        [("Bài báo LinCIR ", B), ("(CVPR 2024) giải bài toán này mà không cần bộ ba gán nhãn: chỉ huấn luyện bằng văn bản.", {})],
        [("Câu hỏi của nhóm: ", B), ("sửa hàm mất mát (A) hoặc sửa nhiễu (B) có giúp LinCIR truy xuất tốt hơn không?", {})]],
        size=20, line=1.05)

    # N3 (old 8): metrics with formulas + evaluation protocol
    s = o[8]
    set_paras(find(s, "TextBox 10"), ["R@K = số truy vấn có ảnh đích trong top-K / tổng số truy vấn",
                                      "AP@K = Σ P@k · rel(k) / min(K, số ảnh đúng); mAP@K lấy trung bình trên các truy vấn"])
    set_paras(find(s, "TextBox 14"), ["Nhóm đánh giá trên FashionIQ val (R@10, R@50) và CIRCO val (mAP@K), "
                                      "kiểm định bootstrap và Wilcoxon theo từng truy vấn"])

    # N4 (old 4): token idea, one line on what is learned
    s = o[4]
    text(s, 0.6, 6.0, 11.2, 0.6, [[("Chỉ học mạng chiếu φ; ", B),
                                   ("bộ mã hóa ảnh và văn bản của CLIP giữ nguyên", {})]], size=20)

    # N5 (old 2): paper summary, bottom line now introduces the two components
    s = o[2]
    set_paras(find(s, "TextBox 14"), [["Gu et al. (NAVER), CVPR 2024, mã nguồn công khai. ",
                                       "Hai thành phần chính: Self-Masking Projection và nhiễu khi huấn luyện"]])

    # N6 (old 5): SMP diagram without the group's changes
    s = o[5]
    fit_picture(s, "Picture 5", ROOT / "report" / "figures" / "lincir_train_base_slide.png", 1.2, 1.75, 10.9, 5.0)

    # N7 (new): noise and modality gap, same layout as the A/B slides
    n7 = frame(prs, "Nhiễu giúp φ dùng được vector ảnh dù hai miền lệch nhau", (-1.8, 4.6, 4.0), (0.06, 3.4, 0.45))
    card(n7, 0.6, 1.9, 6.1, 2.2, "Modality gap của CLIP", ["Vector ảnh và vector văn bản nằm ở hai vùng tách biệt",
                                                          "Khi suy luận, φ nhận vector ảnh dù chỉ học trên văn bản"])
    card(n7, 0.6, 4.3, 6.1, 2.2, "Nhiễu của LinCIR", ["Khi huấn luyện, φ nhận z_c + u · ε",
                                                     "u ~ U(0, 1), ε ~ N(0, I): độ lớn nhiễu đa dạng"])
    stat(n7, 7.3, 2.0, 5.2, "0,82", "khoảng cách giữa tâm vector ảnh và tâm vector văn bản (5.000 cặp COCO)")
    stat(n7, 7.3, 4.3, 5.2, "0,26", "cosine trung bình giữa một ảnh và chú thích của chính nó")

    # N8 (old 6): improvement A, hypothesis instead of the result
    s = o[6]
    set_paras(find(s, "TextBox 13"), ["H1"])
    set_paras(find(s, "TextBox 14"), ["giả thuyết: A tăng R@K, mAP@K so với baseline"])

    # N9 (old 7): improvement B, problem / solution / hypothesis
    s = o[7]
    set_paras(find(s, "TextBox 6"), ["Vấn đề"])
    set_paras(find(s, "TextBox 7"), ["Nhiễu u · ε đều theo mọi hướng", "Không dịch về phía vùng ảnh"])
    set_paras(find(s, "TextBox 9"), ["Giải pháp: nhiễu có hướng"])
    set_paras(find(s, "TextBox 12"), ["độ lớn của Δ, tính trên 5.000 cặp COCO val2017 (khác tập đánh giá)"])
    set_paras(find(s, "TextBox 13"), ["H2"])
    set_paras(find(s, "TextBox 14"), ["giả thuyết: φ nhất quán hơn giữa ảnh và văn bản, nên R@K, mAP@K tăng"])

    # N10 (old 9): code + Kaggle experiments (old 10 is removed)
    s = o[9]
    set_title(s, "Dựng lại mã nguồn và chạy 4 biến thể × 3 seed trên Kaggle")
    set_paras(find(s, "TextBox 13"), ["Thực nghiệm trên Kaggle"])
    set_paras(find(s, "TextBox 14"), ["Baseline, A, B, A+B (H3: A+B có cộng hưởng không)",
                                      "3 seed, 5.000 bước, chỉ khác thành phần cải tiến",
                                      "2 GPU T4; nạp dữ liệu nhanh gấp 3,1 lần"])

    # N11 (old 11): reproduction + paper/code differences
    s = o[11]
    drop(s, "TextBox 8", "TextBox 9", "TextBox 10", "TextBox 11")
    card(s, 4.7, 4.1, 8.0, 2.45, "Bài báo khác mã nguồn (nhóm theo mã)",
         ["φ trong mã không có LayerNorm (bài báo có)", "Mã nạp thêm câu prompt Midjourney",
          "Từ khóa bị che gồm cả mạo từ (a, the)"])

    # N12 (old 12): add the official checkpoint row, highlight B
    s = o[12]
    tbl_shape = find(s, "Table 5")
    tbl = tbl_shape.table
    tr = copy.deepcopy(tbl.rows[1]._tr)
    tbl._tbl.append(tr)
    new = tbl.rows[len(tbl.rows) - 1]
    for c, v in zip(new.cells, ["Checkpoint chính thức (20.000 bước)", "26,35", "46,12", "10,50"]):
        set_paras(c, [v])
    for c in tbl.rows[3].cells:  # B
        c.fill.solid()
        c.fill.fore_color.rgb = TINT
        for r in c.text_frame.paragraphs[0].runs:
            r.font.bold = True
    for c in new.cells:
        for r in c.text_frame.paragraphs[0].runs:
            r.font.color.rgb = RGBColor(0x5B, 0x6B, 0x82)
            r.font.italic = True
    rh = tbl.rows[1].height
    for name in ("Rounded Rectangle 6", "TextBox 7"):
        sh = find(s, name)
        sh.top += rh
    find(s, "Rounded Rectangle 6").height -= rh - Inches(0.1)

    # N13 (old 13): smaller charts + takeaway text
    s = o[13]
    fit_picture(s, "Picture 9", ROOT / "results" / "figures" / "modality_gap_slide.png", 0.6, 1.75, 5.9, 3.35)
    fit_picture(s, "Picture 10", ROOT / "results" / "figures" / "learning_curve_slide.png", 6.8, 1.75, 5.9, 3.35)
    text(s, 0.6, 5.25, 11.2, 1.5, [
        [("Trái: ", B), ("cosine giữa φ(ảnh) và φ(văn bản) của cùng một cặp; B cao nhất (0,84), A thấp nhất (0,64), "
                          "cùng thứ tự với CIRCO mAP@5", {})],
        [("Phải: ", B), ("A thấp hơn baseline ở mọi mốc huấn luyện, nên kết luận không phụ thuộc điểm dừng", {})]],
        size=20, line=1.05)

    # N14 (old 14): conclusion with H1-H3 + retrieval example
    s = o[14]
    set_title(s, "Kết luận: H2 được ủng hộ, H1 và H3 bị bác bỏ")
    drop(s, "Rounded Rectangle 5", "TextBox 6", "Rounded Rectangle 7", "TextBox 8", "TextBox 9")
    q = ROOT / "results" / "figures" / "qualitative_fiq_1.png"
    iw, ih = Image.open(q).size
    w = 7.6
    h = w * ih / iw
    s.shapes.add_picture(str(q), Inches(0.6), Inches(1.75), Inches(w), Inches(h))
    y = 1.75 + h + 0.1
    text(s, 0.6, y, w, 0.8, "Ví dụ FashionIQ: B xếp đúng ảnh đích ở hạng 1; baseline trả về váy ngắn, tối màu",
         size=20, color=NAVY)
    text(s, 0.6, y + 0.95, w, 0.8, [[("Đóng góp: ", B), ("dựng lại mã, tái hiện, hai cải tiến, kiểm định thống kê, "
                                                          "phân tích cơ chế", {})]], size=20)
    card(s, 8.45, 1.75, 4.3, 4.85, "Kết luận", [
        "H2 (B) được ủng hộ: R@10 +0,60, R@50 +1,27; vượt checkpoint 20.000 bước",
        "H1 (A) bị bác bỏ: CIRCO mAP@5 giảm 2,45 điểm",
        "H3 (A+B) bị bác bỏ: kém baseline trên CIRCO",
        "Hạn chế: 5.000 bước, chưa quét λ, α, chưa có CIRR"], size=20)

    # order: old ids, "N7" = new slide; old 10 removed
    order = [1, 3, 8, 4, 2, 5, "N7", 6, 7, 9, 11, 12, 13, 14, 15]
    lst = prs.slides._sldIdLst
    ids = list(lst)
    by_old = {i + 1: e for i, e in enumerate(ids[:15])}
    by_old["N7"] = ids[15]
    rid10 = by_old[10].get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
    for e in ids:
        lst.remove(e)
    for k in order:
        lst.append(by_old[k])
    prs.part.drop_rel(rid10)
    for i, sl in enumerate(prs.slides, start=1):
        for sh in sl.shapes:
            if sh.name == "!!num":
                runs = sh.text_frame.paragraphs[0].runs
                runs[0].text = str(i)
                for r in runs[1:]:
                    r.text = ""
    prs.save(deck)
    print("đã tái cấu trúc:", len(prs.slides), "slide")


if __name__ == "__main__":
    main()
