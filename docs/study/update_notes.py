"""[NHOM8] Rewrite the speaker notes of the (hand-edited) deck in plain language. Slides themselves are untouched.
  python docs/study/update_notes.py [--deck slides/Slides_Nhom8_LinCIR.pptx]"""
import argparse
import copy
import datetime as dt
import shutil
import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from plain_language import SLIDES  # noqa: E402

SPEAKER = {"1": "Tâm", "2": "Tâm", "3": "Tâm", "4": "Tâm", "5": "Tâm", "6": "Tâm", "7": "Tâm", "8": "Tâm", "9": "Tâm", "10": "Tâm", "11": "Tâm", "12": "Tâm", "13": "Tâm", "14": "Tâm", "15": "Tâm", "16": "Tâm"}

ap = argparse.ArgumentParser()
ap.add_argument("--deck", default=str(Path(__file__).resolve().parents[2] / "slides" / "Slides_Nhom8_LinCIR.pptx"))
a = ap.parse_args()
deck = Path(a.deck)
shutil.copy(deck, deck.with_name(f"{deck.stem}_truocKhiSuaGhiChu_{dt.datetime.now():%Y%m%d-%H%M}.pptx"))
prs = Presentation(deck)
body_src = next(sl.notes_slide.notes_placeholder._element for sl in prs.slides
                if sl.has_notes_slide and sl.notes_slide.notes_placeholder is not None)
for i, s in enumerate(prs.slides, start=1):
    if s.notes_slide.notes_text_frame is None:  # slide added without a notes body: copy one
        s.notes_slide.shapes._spTree.append(copy.deepcopy(body_src))
    k = str(i)
    if k in SLIDES:
        say, number = SLIDES[k]
        lines = [f"[{SPEAKER[k]}] {say}"]
        if number != "—":
            lines.append(f"Con số cần nêu: {number}.")
    else:
        lines = [f"[{SPEAKER.get(k, 'Cả nhóm')}] Cảm ơn cô và các bạn; mời đặt câu hỏi.",
                 "Chuyển câu hỏi: phương pháp, cải tiến A → Trung; tái hiện, bảng kết quả → Tùng; độ đo, kiểm định, cải tiến B, cơ chế → Dũng; bài toán, lý thuyết, ví dụ định tính → Nou; mã nguồn, Kaggle, thiết lập → Tâm."]
    s.notes_slide.notes_text_frame.text = "\n".join(lines)
prs.save(deck)
print("đã cập nhật ghi chú", len(prs.slides), "slide")
